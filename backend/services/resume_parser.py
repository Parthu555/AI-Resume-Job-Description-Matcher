import io
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Dict, List, Optional, Tuple

try:
    import fitz  # PyMuPDF
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    import pypdf
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    import docx
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False


COMMON_SECTIONS = [
    "summary",
    "professional summary",
    "objective",
    "experience",
    "work experience",
    "professional experience",
    "employment history",
    "education",
    "academic background",
    "skills",
    "technical skills",
    "core competencies",
    "projects",
    "key projects",
    "certifications",
    "licenses & certifications",
    "publications",
    "awards",
    "internships",
    "volunteer experience"
]


class ResumeParsingError(Exception):
    """Custom exception raised when resume parsing fails with a user-friendly message."""
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


class ResumeParser:
    def __init__(self, max_size_mb: int = 5, min_chars: int = 50):
        self.max_size_bytes = max_size_mb * 1024 * 1024
        self.min_chars = min_chars

    def validate_file(self, filename: str, file_bytes: bytes) -> str:
        """Validates file extension, size, and non-empty content."""
        if not file_bytes or len(file_bytes) == 0:
            raise ResumeParsingError("The uploaded file is completely empty. Please upload a valid resume.")

        if len(file_bytes) > self.max_size_bytes:
            raise ResumeParsingError(
                f"File size exceeds the maximum allowed limit of {self.max_size_bytes // (1024 * 1024)} MB."
            )

        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            return "pdf"
        elif lower_name.endswith(".docx"):
            return "docx"
        else:
            raise ResumeParsingError(
                "Unsupported file format. Please upload a document with .pdf or .docx extension."
            )

    def extract_text_from_pdf(self, file_bytes: bytes) -> str:
        """Extracts text from PDF bytes using PyMuPDF if available, or pypdf fallback."""
        text = ""

        # Attempt 1: PyMuPDF (fitz)
        if PYMUPDF_AVAILABLE:
            try:
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                if doc.is_encrypted:
                    raise ResumeParsingError("The PDF file is password protected or encrypted. Please remove encryption.")
                for page in doc:
                    text += page.get_text("text") + "\n"
                doc.close()
            except ResumeParsingError:
                raise
            except Exception as e:
                # If PyMuPDF fails or file is malformed, attempt fallback
                pass

        # Attempt 2: pypdf (fallback if PyMuPDF wasn't used or yielded empty)
        if not text.strip() and PYPDF_AVAILABLE:
            try:
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                if reader.is_encrypted:
                    raise ResumeParsingError("The PDF file is password protected or encrypted. Please remove encryption.")
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
            except ResumeParsingError:
                raise
            except Exception as e:
                raise ResumeParsingError(f"Failed to read PDF file: {str(e)}")

        clean_text = self.normalize_text(text)
        if len(clean_text) < self.min_chars:
            raise ResumeParsingError(
                "No readable text found in the PDF. The file may be a scanned image or non-searchable document. "
                "Optical Character Recognition (OCR) is required for scanned resumes."
            )
        return clean_text

    def extract_text_from_docx(self, file_bytes: bytes) -> str:
        """Extracts text from DOCX bytes using python-docx or native ZIP/XML parser."""
        text = ""

        # Attempt 1: python-docx if installed
        if PYTHON_DOCX_AVAILABLE:
            try:
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also include table text
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            paragraphs.append(row_text)
                text = "\n".join(paragraphs)
            except Exception:
                pass

        # Attempt 2: Native ZIP XML parser (standard library, zero external dependency)
        if not text.strip():
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                    if "word/document.xml" not in z.namelist():
                        raise ResumeParsingError("Invalid or corrupted DOCX file: word/document.xml missing.")
                    xml_content = z.read("word/document.xml")
                    root = ET.fromstring(xml_content)
                    
                    # Namespaces for WordprocessingML
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    paragraphs = []
                    for p in root.findall(".//w:p", ns):
                        texts = [node.text for node in p.findall(".//w:t", ns) if node.text]
                        if texts:
                            paragraphs.append("".join(texts))
                    text = "\n".join(paragraphs)
            except zipfile.BadZipFile:
                raise ResumeParsingError("The uploaded DOCX file is corrupt or not a valid Word document.")
            except ResumeParsingError:
                raise
            except Exception as e:
                raise ResumeParsingError(f"Failed to extract text from DOCX: {str(e)}")

        clean_text = self.normalize_text(text)
        if len(clean_text) < self.min_chars:
            raise ResumeParsingError(
                "Extracted text from DOCX is too short for meaningful analysis. Please upload a complete resume."
            )
        return clean_text

    def parse_file(self, filename: str, file_bytes: bytes) -> Tuple[str, List[str]]:
        """Main parsing function returning (normalized_text, detected_sections)."""
        file_type = self.validate_file(filename, file_bytes)
        if file_type == "pdf":
            text = self.extract_text_from_pdf(file_bytes)
        else:
            text = self.extract_text_from_docx(file_bytes)

        sections = self.detect_sections(text)
        return text, sections

    def normalize_text(self, text: str) -> str:
        """Cleans and standardizes extracted text while retaining structural breaks."""
        if not text:
            return ""
        # Normalize non-breaking spaces and strange unicode spaces
        text = text.replace("\u00a0", " ").replace("\u200b", "")
        # Standardize bullet characters
        text = re.sub(r"[\u2022\u2023\u25e6\u2043\u2219\uf0b7]", "\n- ", text)
        # Collapse excessive carriage returns
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        # Collapse 3 or more newlines to 2
        text = re.sub(r"\n{3,}", "\n\n", text)
        # Collapse multiple spaces per line
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
        return "\n".join(lines).strip()

    def detect_sections(self, text: str) -> List[str]:
        """Identifies structured resume sections present in the text."""
        detected = []
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        
        # Mapping variations to canonical section names
        canonical_map = {
            "professional summary": "Summary",
            "summary": "Summary",
            "objective": "Summary",
            "work experience": "Experience",
            "professional experience": "Experience",
            "experience": "Experience",
            "employment history": "Experience",
            "technical skills": "Skills",
            "skills": "Skills",
            "core competencies": "Skills",
            "education": "Education",
            "academic background": "Education",
            "projects": "Projects",
            "key projects": "Projects",
            "certifications": "Certifications",
            "licenses & certifications": "Certifications",
            "publications": "Publications",
            "awards": "Awards",
            "internships": "Internships",
            "volunteer experience": "Volunteer"
        }

        for line in lines:
            # Strip markdown/bullet decorations
            clean_line = re.sub(r"^[\#\*\-\•\>\s]+", "", line).strip()
            for key, canonical_name in canonical_map.items():
                pattern = rf"^{re.escape(key)}(?:\s*[:\-]|\s*$|\s+[–—])"
                if re.match(pattern, clean_line, re.IGNORECASE):
                    if canonical_name not in detected:
                        detected.append(canonical_name)
                    break
        return detected
